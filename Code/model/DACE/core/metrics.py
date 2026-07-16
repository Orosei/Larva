import os
import time

import cv2
import numpy as np
import pandas as pd
from openpyxl.styles import Border, Font, PatternFill, Side


class TrajCalc:

    def __init__(self, label, band_r, tdice_r):
        self.label = label.astype(np.uint8)
        self.band_r = band_r
        self.tdice_r = tdice_r
        self.eps = 1e-6
        self.skel = self._skel()
        self.dist_field = self._dist_field()

    # ── init helpers ─────────────────────────────────────────────
    def vis(self, pred, skel_band=None, save_path=None):
        base = cv2.cvtColor((self.label > 0).astype(np.uint8) * 128, cv2.COLOR_GRAY2BGR)
        pred_mask = (pred > 0).astype(np.uint8)
        base[pred_mask > 0] = (0, 0, 255)
        if skel_band is not None:
            skel_mask = (skel_band > 0).astype(np.uint8)
            base[skel_mask > 0] = (0, 255, 0)
        if save_path:
            cv2.imwrite(save_path, base)
        cv2.imshow("TraceMetrics.vis", base)
        cv2.waitKey(0)
        cv2.destroyAllWindows()

    def _skel(self):
        mask = (self.label > 0.5).astype(np.uint8) * 255
        skel = cv2.ximgproc.thinning(mask, cv2.ximgproc.THINNING_GUOHALL)
        return (skel > 0).astype(np.uint8)

    def _dist_field(self, window_size=8, eps=1e-6):
        kd = cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (3, 3))
        ld = cv2.dilate(self.label, kd, iterations=1)
        df = cv2.distanceTransform(ld, cv2.DIST_L2, 5)
        lm = cv2.dilate(df, np.ones((window_size, window_size), np.uint8), iterations=1)
        dn = df / (lm + eps)
        dn[self.skel > 0] = 1.0
        return dn.astype(np.float32)

    # ── prepare ──────────────────────────────────────────────────

    def _prepare(self, t):
        x = t[:, 0].astype(np.int32)
        y = t[:, 1].astype(np.int32)
        h, w = self.label.shape
        ok = (x >= 0) & (x < w) & (y >= 0) & (y < h)
        to = t[ok]
        pts = np.array(t, dtype=np.int32).reshape(-1, 1, 2)
        traj_p1 = np.zeros_like(self.label, dtype=np.uint8)
        cv2.polylines(traj_p1, [pts], False, 1, 1)
        band = np.zeros_like(self.label, dtype=np.uint8)
        cv2.polylines(band, [pts], False, 1, self.band_r * 2)
        skel_band = (self.skel > 0).astype(np.uint8) & band

        # self.vis(self.label, traj_p1)
        # self.vis(self.skel, skel_band)

        return to, traj_p1, skel_band

    def _prepare_img(self, pred):
        pred = (pred > 0).astype(np.uint8)
        kr = cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (2 * self.band_r + 1, 2 * self.band_r + 1))
        pred_band = cv2.dilate(pred, kr)
        skel_band = (self.skel > 0).astype(np.uint8) & pred_band
        # self.vis(self.label, pred)
        # self.vis(self.skel, skel_band)
        return pred, skel_band

    # ── eval ─────────────────────────────────────────────────────

    def eval(self, t, time_ms):
        to, traj_p1, skel_band = self._prepare(t)
        m = {
            "IR": self.inside_ratio(to),
            "MD": self.mean_dist(to),
            "TD": self.traj_directness(to),
            "Tdice": self.tdice(traj_p1, skel_band),
            "f1": self.prf(traj_p1, skel_band),
            "t": time_ms,
        }
        return {k: round(float(v), 4) for k, v in m.items()}

    def eval_img(self, pred, time_ms):
        pred, skel_band = self._prepare_img(pred)
        m = {
            "IR": self.inside_ratio_img(pred),
            "MD": self.mean_dist_img(pred),
            "TD": self.traj_directness_img(pred),
            "Tdice": self.tdice(pred, skel_band),
            "f1": self.prf(pred, skel_band),
            "t": time_ms,
        }
        return {k: round(float(v), 4) for k, v in m.items()}

    # ── trajectory metrics ───────────────────────────────────────

    def inside_ratio(self, t):
        x = t[:, 0].astype(np.int32)
        y = t[:, 1].astype(np.int32)
        return np.mean(self.label[y, x] == 1)

    def mean_dist(self, t):
        x = t[:, 0].astype(np.int32)
        y = t[:, 1].astype(np.int32)
        return np.mean(self.dist_field[y, x])

    def _segment(self, t, n):
        segs = []
        for i in range(0, len(t), n):
            seg = t[i : i + n]
            if len(seg) < n and segs:
                segs[-1] = t[i - n :]  # merge tail
            else:
                segs.append(seg)
        return segs

    def _directness(self, seg):
        d = seg[1:] - seg[:-1]
        return np.linalg.norm(seg[-1] - seg[0]) / (np.sum(np.maximum(np.linalg.norm(d, axis=-1), 1.0)) + self.eps)

    def traj_directness(self, t, n=10):
        segs = self._segment(t, n)
        scores = [self._directness(s) for s in segs]
        return float(np.mean(scores)) if scores else 0.0

    def tdice(self, traj_p1, skel_band):
        k_p = np.sum(traj_p1 & (self.label > 0)) / (np.sum(traj_p1) + self.eps)
        kr = cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (2 * self.tdice_r + 1, 2 * self.tdice_r + 1))
        pd = cv2.dilate(traj_p1, kr).astype(np.uint8)
        k_s = np.sum(skel_band & (pd > 0)) / (np.sum(skel_band) + self.eps)
        return 2 * k_p * k_s / (k_p + k_s + self.eps)

    def prf(self, traj_p1, skel_band):
        tp = np.sum((skel_band & traj_p1) > 0)
        fp = np.sum(traj_p1) - tp
        fn = np.sum(skel_band) - tp
        p = tp / (tp + fp + self.eps)
        r = tp / (tp + fn + self.eps)
        return 2 * p * r / (p + r + self.eps)

    # ── image metrics ────────────────────────────────────────────

    def inside_ratio_img(self, pred):
        return float(np.sum(pred & (self.label > 0)) / (np.sum(pred) + self.eps))

    def mean_dist_img(self, pred):
        return float(np.sum(self.dist_field[pred > 0]) / (np.sum(pred) + self.eps))

    def traj_directness_img(self, pred):
        ys, xs = np.where(pred > 0)
        if len(xs) < 2:
            return 0.0
        pts = set(zip(xs, ys))
        start = min(
            pts,
            key=lambda p: sum(
                1 for dx in (-1, 0, 1) for dy in (-1, 0, 1) if (dx or dy) and (p[0] + dx, p[1] + dy) in pts
            ),
        )
        ordered, visited, cur = [start], {start}, start
        while True:
            x, y = cur
            nbrs = [
                (x + dx, y + dy)
                for dx in (-1, 0, 1)
                for dy in (-1, 0, 1)
                if (dx or dy) and (x + dx, y + dy) in pts and (x + dx, y + dy) not in visited
            ]
            if not nbrs:
                break
            nxt = nbrs[0]
            ordered.append(nxt)
            visited.add(nxt)
            cur = nxt
        return self.traj_directness(np.array(ordered, dtype=np.float32))


class TraceMetrics:
    _METRICS = {
        "MD": {"label": "Mean Distance", "important": True},
        "IR": {"label": "Inside Ratio", "important": True},
        "TD": {"label": "Trajectory Directness", "important": True},
        "Tdice": {"label": "Trajectory Dice", "important": True},
        "f1": {"label": "Trajectory F1", "important": False},
        "t": {"label": "Time (ms)", "important": False},
    }

    OUTPUT_ORDER = list(_METRICS.keys())
    IMPORTANT_KEYS = {k for k, v in _METRICS.items() if v["important"]}

    def __init__(
        self,
        label=None,
        base_name="all",
        band_r=4,
        tdice_r=3,
    ):
        self.base_name = base_name
        self.history = []
        self.trajs = []
        self._names = {k: v["label"] for k, v in self._METRICS.items()}
        if label is not None:
            self.calc = TrajCalc(label, band_r, tdice_r)
            self.h, self.w = label.shape[:2]
        else:
            self.calc = None
            self.h = self.w = 0

    # — internals —————————————————————————————————————————————————
    def _avg(self):
        if not self.history:
            return {}
        sums = {}
        counts = {}
        for m in self.history:
            for k, v in m.items():
                sums[k] = sums.get(k, 0.0) + v
                counts[k] = counts.get(k, 0) + 1
        return {k: round(sums[k] / counts[k], 4) for k in sums}

    @staticmethod
    def _fmt(k, v):
        G, R = "\033[92m", "\033[0m"
        return f"{G}{v:>10.4f}{R}" if k in TraceMetrics.IMPORTANT_KEYS else f"{v:>10.4f}"

    def _print_one(self, metrics_dict, title):
        G, R = "\033[92m", "\033[0m"
        print(f"\n{'='*50}\n📊 {title}\n{'-'*50}")
        for k in self.OUTPUT_ORDER:
            name = self._names.get(k, k)
            print(f"  {name:<22} : {self._fmt(k, metrics_dict[k])}")
        imp = [metrics_dict[k] for k in self.IMPORTANT_KEYS]
        print(f"{'-'*50}\n  {G}OVERALL SCORE{R}       : {G}{round(np.mean(imp), 4):>10.4f}{R}\n{'='*50}\n")

    def _print_all(self):
        n, w, cw = len(self.history), 24, 12
        bar = "=" * (w + n * cw)
        print(f"\n{bar}")
        print(f"{'Metric':<24}" + "".join(f"| {'Score ' + str(i + 1):>10}" for i in range(n)))
        print(f"{'-' * (w + n * cw)}")
        for k in self.OUTPUT_ORDER:
            name = self._names.get(k, k)
            vals = [f"| {self._fmt(k, h[k])}" for h in self.history]
            print(f"{name:<24}{''.join(vals)}")
        print(f"{bar}\n")

    def _save_xls(self, save_dir):
        os.makedirs(save_dir, exist_ok=True)
        rows = [{"id": i + 1, **d} for i, d in enumerate(self.history)]
        rows.append({"id": "AVG", **self._avg()})
        df = pd.DataFrame(rows)
        cols = ["id"] + [k for k in self.OUTPUT_ORDER if k in df.columns]
        df = df[cols]

        path = os.path.join(save_dir, f"{self.base_name}_metrics.xlsx")
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            df.to_excel(writer, index=False)
            ws = writer.sheets["Sheet1"]

            bd = Border(
                left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin")
            )
            rf = Font(color="FF0000")
            lf = PatternFill(start_color="DAEEF3", end_color="DAEEF3", fill_type="solid")

            for row in ws.iter_rows():
                for cell in row:
                    cell.border = bd

            for col_idx, col_name in enumerate(cols, start=1):
                if col_name in self.IMPORTANT_KEYS:
                    for row in range(2, ws.max_row + 1):
                        ws.cell(row=row, column=col_idx).font = rf

            for cell in ws[ws.max_row]:
                cell.fill = lf

    def _save_traj(self, save_dir):
        os.makedirs(save_dir, exist_ok=True)
        path = os.path.join(save_dir, f"{self.base_name}_traj.csv")

        data = []
        for tid, traj in enumerate(self.trajs, 1):
            for x, y in traj:
                data.append([tid, round(float(x), 1), round(float(y), 1)])

        df = pd.DataFrame(data, columns=["traj", "x", "y"])
        meta = f"# h={self.h},w={self.w}\n"
        csv_text = df.to_csv(index=False, lineterminator="\n")

        with open(path, "w", encoding="utf-8", newline="") as f:
            f.write(meta + csv_text)

    # — public ————————————————————————————————————————————————————
    def eval(self, traj, t0=None):
        time_ms = (time.time() - t0) * 1000 if t0 is not None else None
        traj_arr = np.array(traj)
        metrics = self.calc.eval(traj_arr, time_ms)
        self.history.append(metrics)
        self.trajs.append(traj_arr)

    def append(self, other):
        avg_data = other._avg()
        self.history.append(avg_data)

    def last_str(self):
        last = self.history[-1]
        parts = []
        for key in self.OUTPUT_ORDER:
            if key in last:
                parts.append(f"{key}:{last[key]:.4f}")
        return "  ".join(parts)

    def report(self, save_dir, show_avg=False):
        self._print_all()
        if show_avg:
            self._print_one(self._avg(), "AVERAGE METRICS")
        self._save_xls(save_dir)
        if self.trajs:
            self._save_traj(save_dir)


class CompareMetrics(TraceMetrics):

    def __init__(
        self,
        label,
        pred,
        base_name="all",
    ):
        super().__init__(label, base_name)
        self.pred = pred

    def eval(self, traj, t0=None):
        time_ms = (time.time() - t0) * 1000 if t0 is not None else None
        pred1 = self._pred_band(self.pred, traj)
        metrics = self.calc.eval_img(pred1, time_ms)
        self.history.append(metrics)
        return metrics

    def _pred_band(self, pred, traj, dilate_r=5):
        kernel = cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (2 * dilate_r + 1, 2 * dilate_r + 1))
        traj_d = cv2.dilate(traj, kernel)
        mask = cv2.bitwise_and(pred, pred, mask=traj_d)
        mask = (mask > 0).astype(np.uint8) * 255
        try:
            return cv2.ximgproc.thinning(mask, cv2.ximgproc.THINNING_GUOHALL)
        except AttributeError:
            from skimage.morphology import skeletonize

            return (skeletonize(mask > 0) * 255).astype(np.uint8)
